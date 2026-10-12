"""Offline regression tests. Run: python -m unittest discover -s tests"""

import json
import os
import re
import sys
import tempfile
import unittest

# Never touch the real key files, monthly counters or search cache of whoever runs the tests.
os.environ["HOME"] = tempfile.mkdtemp(prefix="ngo-tests-")
for _k in ("OLA_MAPS_API_KEY", "GOOGLE_MAPS_API_KEY"):
    os.environ.pop(_k, None)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
FIX = os.path.join(ROOT, "tests", "fixtures")

from scraper import districts, enrich  # noqa: E402
from scraper.classify import categories, looks_governmental  # noqa: E402
from scraper.models import Org, merge  # noqa: E402
from scraper.sources import awbi, osm, websearch  # noqa: E402
from scraper import lookup_page, people  # noqa: E402
from scraper.text import clean, name_key  # noqa: E402

LONG_DASH = re.compile("[" + "".join(chr(c) for c in (0x2012, 0x2013, 0x2014, 0x2015)) + "]")


def fixture(name, mode="r"):
    with open(os.path.join(FIX, name), mode) as fh:
        return fh.read()


class FakeResponse:
    def __init__(self, text="", status=200, content=None, url="", ctype="text/html"):
        self.text = text
        self.status_code = status
        self.content = content if content is not None else text.encode()
        self.url = url
        self.headers = {"Content-Type": ctype}

    def json(self):
        return json.loads(self.text)


class FakeHttp:
    """Routes URLs to fixtures; anything unknown is a 404."""

    def __init__(self, routes):
        self.routes = routes
        self.calls = []
        self.params = []

    def _find(self, url):
        self.calls.append(url)
        for key, resp in self.routes.items():
            if key in url:
                return resp() if callable(resp) else resp
        return FakeResponse("not found", 404, url=url)

    def get(self, url, **kw):
        self.params.append(kw.get("params") or {})
        return self._find(url)

    def post(self, url, **kw):
        return self._find(url)


# ------------------------------------------------------------------ AWBI

EXPECTED_KA = {
    "KA001/1965": ("Mysore Panjrapole Society", "Mysuru"),
    "KA004/1993": ("Compassion Unlimited Plus Action", "Bengaluru Urban"),
    "KA017/1999": ("Wildlife Rescue & Rehabilitation Centre", "Bengaluru Urban"),
    "KA021/2000": ("Animal Care Trust", "Dakshina Kannada"),
    "KA040/2005": ("Swapna Samrakshana Samiti", "Uttara Kannada"),
    "KA067/2018": ("Charlies Animal Rescue Centre (CARE)", "Bengaluru Urban"),
    "KA073/2019": ("Captain's Animal Care Trust", "Udupi"),
    "KA074/2020": ("Shri Umamaheswara Seva Trust", "Shivamogga"),
    "KA077/2021": ("Punyakoti Gorakshana Vedika", "Shivamogga"),
    "KA079/2023": ("Shri Veeranjaneya Religious & Charitable Trust", "Uttara Kannada"),
    "KA086/2025": ("Srinivasa Taarini Charitable Trust", "Bengaluru Urban"),
}


class TestAwbi(unittest.TestCase):
    def check_list(self, pdf_name, expect_method):
        rows, method, counts = awbi.parse_pdf(os.path.join(FIX, pdf_name))
        self.assertEqual(method, expect_method, counts)
        ka = awbi.karnataka_rows(rows)
        got = {r["code"]: r for r in ka}
        self.assertEqual(set(got), set(EXPECTED_KA), "Karnataka rows differ")
        for code, (name, district) in EXPECTED_KA.items():
            self.assertEqual(got[code]["name"], name)
            org = awbi.to_orgs([got[code]], "recognised", "")[0]
            self.assertEqual(org.district, district, code)
        # KA008 has a Karnataka-style code but its state is Delhi.
        self.assertNotIn("KA008/1991", got)
        self.assertEqual(got["KA004/1993"]["valid_upto"], "18.07.2028")
        self.assertTrue(got["KA004/1993"]["address"].startswith("No.18/1"))
        self.assertTrue(got["KA004/1993"]["address"].endswith("560008"))

    def test_ruled_table_pdf(self):
        self.check_list("awbi_recognised_ruled.pdf", "table")

    def test_plain_pdf_without_gridlines(self):
        self.check_list("awbi_recognised_plain.pdf", "layout")

    def test_centred_headers_without_gridlines(self):
        self.check_list("awbi_recognised_centred.pdf", "layout")

    def test_abc_list_skips_municipal_bodies(self):
        rows, _, _ = awbi.parse_pdf(os.path.join(FIX, "awbi_abc_ruled.pdf"))
        orgs = awbi.to_orgs(awbi.karnataka_rows(rows), "abc", "")
        names = sorted(o.name for o in orgs)
        self.assertEqual(names, ["Captain's Animal Care Trust",
                                 "Spandana Animal Welfare & Prakruthi Seva Society"])
        codes = {o.name: o.awbi_code for o in orgs}
        self.assertEqual(codes["Captain's Animal Care Trust"], "KA073/2019")

    def test_lapsed_validity_note(self):
        rows, _, _ = awbi.parse_pdf(os.path.join(FIX, "awbi_recognised_ruled.pdf"))
        row = [r for r in rows if r["code"] == "KA077/2021"][0]
        org = awbi.to_orgs([row], "recognised", "")[0]
        self.assertTrue(any("has passed" in n for n in org.notes))

    def test_fetch_uses_local_pdf_and_survives_bad_urls(self):
        http = FakeHttp({})  # every URL 404s
        orgs, info, _ = awbi.fetch(http, "recognised",
                                   [os.path.join(FIX, "awbi_recognised_ruled.pdf")], None, lambda m: None,
                                   states=["Karnataka"])
        self.assertEqual(len(orgs), len(EXPECTED_KA))
        everything, _, _ = awbi.fetch(http, "recognised",
                                      [os.path.join(FIX, "awbi_recognised_ruled.pdf")], None, lambda m: None)
        self.assertEqual(len(everything), 15)  # every state
        by_name = {o.name: o for o in everything}
        self.assertEqual(by_name["Elephant Welfare Association"].state, "Kerala")
        self.assertEqual(by_name["Indian Vety. Association"].state, "Delhi")
        orgs, info, _ = awbi.fetch(http, "recognised", [], None, lambda m: None)
        self.assertEqual(orgs, [])
        self.assertTrue(info["tried"])

    def test_fetch_downloads_newest_pdf(self):
        pdf = fixture("awbi_recognised_ruled.pdf", "rb")
        http = FakeHttp({"177608352345": FakeResponse(content=pdf, ctype="application/pdf")})
        orgs, info, _ = awbi.fetch(http, "recognised", [], None, lambda m: None, states=["Karnataka"])
        self.assertEqual(len(orgs), len(EXPECTED_KA))
        self.assertIn("177608352345", info["used"])


# ------------------------------------------------------------- districts

class TestDistricts(unittest.TestCase):
    def test_aliases(self):
        cases = {
            "Bangalore": "Bengaluru Urban", "Banagalore": "Bengaluru Urban",
            "Dist-Shimoga": "Shivamogga", "Mangalore": "Dakshina Kannada",
            "Sadashivgad": "Uttara Kannada", "Hubli": "Dharwad", "Gulbarga": "Kalaburagi",
            "UDUPI": "Udupi", "Madikeri": "Kodagu", "T.Narasipur": "Mysuru",
            "Chickballapur": "Chikkaballapura", "Hospet": "Vijayanagara",
            "Ramanagara": "Bengaluru South (Ramanagara)", "Ilkal": "Bagalkote",
            "Thrissur": "", "": "",
        }
        for raw, want in cases.items():
            self.assertEqual(districts.normalise(raw), want, raw)

    def test_guess_inside_address(self):
        self.assertEqual(districts.guess_from_text("24, Ananth Nagar, Manipal, 576104"), "Udupi")

    def test_31_districts_with_search_towns(self):
        self.assertEqual(len(districts.DISTRICTS), 31)
        self.assertEqual(set(districts.SEARCH_TOWN), set(districts.DISTRICTS))


# ------------------------------------------------------------ classifying

class TestClassify(unittest.TestCase):
    def test_categories(self):
        self.assertEqual(categories("Charlies Animal Rescue Centre"), ["animal"])
        self.assertEqual(categories("Nature Conservation Foundation"), ["wildlife"])
        self.assertEqual(categories("Punyakoti Gorakshana Vedika"), ["cattle"])
        self.assertIn("wildlife", categories("Wildlife Rescue & Rehabilitation Centre"))
        self.assertEqual(categories("Mango Tree Education Trust"), [])
        self.assertEqual(categories("Tobago Cargo Services"), [])

    def test_government_detection(self):
        self.assertTrue(looks_governmental("Bruhat Bengaluru Mahanagara Palike"))
        self.assertFalse(looks_governmental("Compassion Unlimited Plus Action"))


# --------------------------------------------------------- text + merging

class TestTextAndMerge(unittest.TestCase):
    def test_clean_removes_long_dashes(self):
        self.assertEqual(clean("Guntur " + chr(0x2013) + " 522003 " + chr(0x2014) + " AP"), "Guntur - 522003 - AP")

    def test_name_key_ignores_filler(self):
        self.assertEqual(name_key("The Animal Care Trust"), name_key("Animal Care Trust (R)"))

    def test_merge_prefers_official_record(self):
        a = Org(name="Charlies Animal Rescue Centre (CARE)", source="AWBI recognised list",
                confidence="high", city="Bengaluru", awbi_code="KA067/2018")
        b = Org(name="Charlies Animal Rescue Centre", source="OpenStreetMap", confidence="medium",
                city="Bengaluru", phones=["+91 98450 00000"], lat="12.9", lon="77.6")
        c = Org(name="Charlies Animal Rescue Centre", source="Web search", district="Bengaluru Urban",
                website="https://charlies.example.org")
        out = merge([c, b, a])
        self.assertEqual(len(out), 1)
        o = out[0]
        self.assertEqual(o.name, "Charlies Animal Rescue Centre (CARE)")
        self.assertEqual(o.confidence, "high")
        self.assertEqual(set(o.sources), {"AWBI recognised list", "OpenStreetMap", "Web search"})
        self.assertEqual(o.phones, ["+91 98450 00000"])
        self.assertEqual(o.website, "https://charlies.example.org")

    def test_merge_keeps_same_name_in_different_districts(self):
        a = Org(name="People for Animals", source="AWBI recognised list", confidence="high",
                city="Mysore", awbi_code="KA016/1999")
        b = Org(name="People for Animals", source="AWBI recognised list", confidence="high",
                city="Hubli", awbi_code="KA026/2002")
        self.assertEqual(len(merge([a, b])), 2)


# ------------------------------------------------------------------- OSM

class TestOsm(unittest.TestCase):
    def test_parse(self):
        orgs = osm.parse(json.loads(fixture("osm_sample.json")))
        names = sorted(o.name for o in orgs)
        self.assertEqual(names, ["Charlies Animal Rescue Centre", "Mysuru Wildlife Trust",
                                 "Udupi Stray Care Shelter"])
        udupi = [o for o in orgs if o.name.startswith("Udupi")][0]
        self.assertEqual(udupi.district, "Udupi")
        self.assertEqual(udupi.phones, ["+91 98450 11111"])
        self.assertEqual(udupi.lat, "13.3409")
        mys = [o for o in orgs if o.name.startswith("Mysuru")][0]
        self.assertEqual((mys.lat, mys.district), ("12.3", "Mysuru"))

    def test_fetch_falls_back_to_second_endpoint(self):
        http = FakeHttp({"overpass-api.de": FakeResponse("busy", 504),
                         "kumi.systems": FakeResponse(fixture("osm_sample.json"))})
        orgs, info = osm.fetch(http, lambda m: None)
        self.assertEqual(len(orgs), 3)
        self.assertIn("kumi", info["endpoint"])


# ------------------------------------------------------------ web search

class TestWebSearch(unittest.TestCase):
    def test_ddg_parse_and_filter(self):
        hits = websearch.parse_ddg(fixture("ddg_results.html"))
        self.assertEqual(len(hits), 7)
        self.assertEqual(hits[0]["url"], "https://www.cupabangalore.org/")
        orgs = websearch.candidates(hits, "Bengaluru Urban")
        names = {o.name: o for o in orgs}
        self.assertEqual(set(names), {"Compassion Unlimited Plus Action",
                                      "Udupi Animal Rescue Group",
                                      "Nature Conservation Foundation"})
        self.assertEqual(names["Udupi Animal Rescue Group"].district, "Udupi")
        self.assertEqual(names["Udupi Animal Rescue Group"].social,
                         ["https://www.facebook.com/udupianimalrescue/"])
        self.assertEqual(names["Udupi Animal Rescue Group"].website, "")
        self.assertEqual(names["Nature Conservation Foundation"].district, "Mysuru")
        self.assertTrue(all(o.confidence == "low" for o in orgs))

    def test_ddg_captcha_raises(self):
        with self.assertRaises(websearch.Blocked):
            websearch.parse_ddg(fixture("ddg_captcha.html"))

    def test_bing_parse_unwraps_links(self):
        hits = websearch.parse_bing(fixture("bing_results.html"))
        self.assertEqual(hits[0]["url"], "https://captainsanimalcare.example.org/home")

    def test_searcher_rests_blocked_engine_and_uses_next(self):
        http = FakeHttp({"duckduckgo": FakeResponse(fixture("ddg_captcha.html")),
                         "bing.com": FakeResponse(fixture("bing_results.html"))})
        s = websearch.Searcher(http, log=lambda m: None, engines=["ddg", "bing"])
        hits = s.search("anything")
        self.assertEqual(len(hits), 2)
        self.assertIn("ddg", s.cool)
        self.assertNotIn("bing", s.cool)

    def test_searcher_rotates_across_engines(self):
        http = FakeHttp({"lite.duckduckgo": FakeResponse(fixture("ddg_lite_results.html")),
                         "mojeek": FakeResponse(fixture("mojeek_results.html")),
                         "brave": FakeResponse(fixture("brave_results.html"))})
        s = websearch.Searcher(http, log=lambda m: None, engines=["ddg_lite", "mojeek", "brave"])
        first = [s.search("q{}".format(i))[0]["title"] for i in range(3)]
        self.assertEqual(first, ["Asare Animal Trust | Surathkal | Facebook",
                                 "Kalaburagi Paws Rescue Trust", "Hassan Strays Rescue (@hassanstrays)"])

    def test_searcher_gives_up_when_everything_is_blocked(self):
        http = FakeHttp({"duckduckgo": FakeResponse(fixture("ddg_captcha.html"))})
        s = websearch.Searcher(http, log=lambda m: None, engines=["ddg"], max_wait=0)
        with self.assertRaises(websearch.Blocked):
            s.search("anything")

    def test_search_cache_avoids_repeat_requests(self):
        from scraper.net import Cache
        http = FakeHttp({"bing.com": FakeResponse(fixture("bing_results.html"))})
        cache = Cache(folder=tempfile.mkdtemp())
        s = websearch.Searcher(http, log=lambda m: None, engines=["bing"], cache=cache)
        s.search("same query")
        s.search("same query")
        self.assertEqual(len(http.calls), 1)

    def test_new_engine_parsers(self):
        lite = websearch.parse_ddg_lite(fixture("ddg_lite_results.html"))
        self.assertEqual(lite[0]["url"], "https://www.facebook.com/asareanimaltrust/")
        self.assertIn("Surathkal", lite[0]["snippet"])
        moj = websearch.parse_mojeek(fixture("mojeek_results.html"))
        self.assertEqual(moj[0]["url"], "https://kalaburagipaws.example.org/")
        brave = websearch.parse_brave(fixture("brave_results.html"))
        self.assertEqual(brave[0]["url"], "https://www.instagram.com/hassanstrays/")
        self.assertEqual(len(brave), 1)

    def test_find_website(self):
        http = FakeHttp({"bing.com": FakeResponse(fixture("bing_results.html"))})
        s = websearch.Searcher(http, log=lambda m: None, engines=["bing"])
        org = Org(name="Captain's Animal Care Trust", source="AWBI", confidence="high", city="Udupi")
        self.assertTrue(websearch.find_website(s, org))
        self.assertEqual(org.website, "https://captainsanimalcare.example.org")
        self.assertEqual(org.social, ["https://www.instagram.com/captainsanimalcare/"])

    def test_name_from_title(self):
        self.assertEqual(websearch.name_from_title("Udupi Animal Rescue Group | Facebook"),
                         "Udupi Animal Rescue Group")
        self.assertEqual(websearch.name_from_title("Paws Trust (@pawstrust) • Instagram photos"),
                         "Paws Trust")


# ------------------------------------------------------------- enrichment

class TestEnrich(unittest.TestCase):
    def test_extract_and_follow_contact_page(self):
        http = FakeHttp({
            "/contact-us": FakeResponse(fixture("site_contact.html"),
                                        url="https://udupistraycare.example.org/contact-us"),
            "/about-us": FakeResponse(fixture("site_about.html"),
                                      url="https://udupistraycare.example.org/about-us"),
            "/our-team": FakeResponse(fixture("site_team.html"),
                                      url="https://udupistraycare.example.org/our-team"),
            "udupistraycare.example.org": FakeResponse(fixture("site_home.html"),
                                                       url="https://udupistraycare.example.org/"),
        })
        org = Org(name="Udupi Stray Care Shelter", source="OpenStreetMap", confidence="medium",
                  website="https://udupistraycare.example.org")
        enrich.enrich(http, org, lambda m: None)
        self.assertEqual(org.emails, ["hello@udupistraycare.example.org",
                                      "adopt@udupistraycare.example.org"])
        self.assertIn("+91 98450 11111", org.phones)
        self.assertIn("0820 2571234", org.phones)
        self.assertIn("+91 98450 22222", org.phones)
        self.assertEqual(org.social, ["https://www.instagram.com/udupistraycare"])
        self.assertFalse(any("tracker" in e for e in org.emails))
        got = {p["name"]: p for p in org.people}
        self.assertEqual(set(got), {"Rekha Shenoy", "Dr. K. S. Bhat"})
        self.assertEqual(got["Rekha Shenoy"]["role"], "Founder & Managing Trustee")
        self.assertEqual(got["Dr. K. S. Bhat"]["role"], "Trustee")
        self.assertTrue(got["Dr. K. S. Bhat"]["source"].endswith("/our-team"))

    def test_phone_normalising(self):
        self.assertEqual(enrich.norm_phone("+91-98450 12345"), "+91 98450 12345")
        self.assertEqual(enrich.norm_phone("098450 12345"), "+91 98450 12345")
        self.assertEqual(enrich.norm_phone("12345"), "")


# ---------------------------------------------------------------- people

class TestPeople(unittest.TestCase):
    def names(self, text, org=""):
        return [(p["name"], p["role"]) for p in people.from_snippet(text, "u", org)]

    def test_common_phrasings(self):
        self.assertEqual(self.names("Founder: Dr. K. S. Ramesh"), [("Dr. K. S. Ramesh", "Founder")])
        self.assertEqual(self.names("Sudha Narayanan, Managing Trustee"),
                         [("Sudha Narayanan", "Managing Trustee")])
        self.assertEqual(self.names("The shelter was founded by Priya Menon in 2010."),
                         [("Priya Menon", "Founder")])
        self.assertEqual(self.names("Managing Trustee - Anand Kumar Read More"),
                         [("Anand Kumar", "Managing Trustee")])

    def test_rejects_non_names(self):
        self.assertEqual(self.names("Captain's Animal Care Trust, Founder", "Captain's Animal Care Trust"), [])
        self.assertEqual(self.names("Copyright 2024 Animal Care Trust. All rights reserved"), [])
        self.assertEqual(self.names("Meet the team Read More"), [])
        self.assertEqual(self.names("our founder said the shelter is full"), [])

    def test_combine_roles(self):
        self.assertEqual(people.combine_roles("Founder", "Founder & Managing Trustee"),
                         "Founder & Managing Trustee")
        self.assertEqual(people.combine_roles("Trustee", "Managing Trustee"), "Managing Trustee")
        self.assertEqual(people.combine_roles("President", "Founder"), "President & Founder")

    def test_card_layout(self):
        found = people.from_lines(["Our Team", "Arjun Shetty", "Founder", "Meera Rao", "Secretary", "Donate"], "u")
        self.assertEqual([(p["name"], p["role"]) for p in found],
                         [("Arjun Shetty", "Founder"), ("Meera Rao", "Secretary")])

    def test_search_fallback_uses_only_snippets_about_the_org(self):
        http = FakeHttp({"duckduckgo": FakeResponse(fixture("ddg_people.html"))})
        s = websearch.Searcher(http, log=lambda m: None, engines=["ddg"])
        org = Org(name="Compassion Unlimited Plus Action", source="AWBI recognised list",
                  confidence="high", city="Bangalore")
        websearch.find_people(s, org)
        self.assertEqual([p["name"] for p in org.people], ["Suparna Ganguly"])
        self.assertEqual(org.people[0]["source"], "https://news.example.org/cupa-story")

    def test_people_survive_merge(self):
        a = Org(name="Animal Care Trust", source="AWBI recognised list", confidence="high", city="Mangalore")
        b = Org(name="Animal Care Trust", source="Web search", district="Dakshina Kannada",
                people=[{"key": "sumamathur", "name": "Suma Mathur", "role": "Founder", "source": "u"}])
        out = merge([b, a])
        self.assertEqual(len(out), 1)
        self.assertEqual(out[0].people[0]["name"], "Suma Mathur")


# --------------------------------------------------------- google places

class TestGooglePlaces(unittest.TestCase):
    def setUp(self):
        from scraper.sources import google_places as gp
        self.gp = gp
        self.page = json.loads(fixture("places_page.json"))

    def test_keep_filters_listings(self):
        kept = [p["displayName"]["text"] for p in self.page["places"] if self.gp.keep(p)]
        self.assertEqual(kept, ["Asare Animal Trust", "Ahimsa Animal Care Trust",
                                "CUPA - Compassion Unlimited Plus Action", "Sri Krishna Gaushala"])

    def test_lead_district_from_address(self):
        lead = self.gp.to_lead(self.page["places"][0], "q", "Udupi")
        self.assertEqual(lead.district, "Dakshina Kannada")
        self.assertEqual(lead.maps_link, "https://www.google.com/maps/place/?q=place_id:ChIJasare")

    def test_query_plan_covers_every_district_and_town(self):
        plan = self.gp.query_plan()
        covered = {d for _, d in plan}
        self.assertEqual(covered, set(districts.DISTRICTS))
        self.assertTrue(any("Surathkal" in q for q, _ in plan))
        self.assertTrue(any("Karkala" in q for q, _ in plan))
        self.assertLess(len(plan) * 2, 4500)  # stays inside the free monthly calls

    def test_sweep_respects_monthly_cap(self):
        calls = []

        class Http:
            def post(self_, url, **kw):
                calls.append(kw["json"]["textQuery"])
                return FakeResponse(fixture("places_page.json"))
        usage_file = os.path.join(tempfile.mkdtemp(), "usage.json")
        orig = self.gp.USAGE_FILE
        self.gp.USAGE_FILE = usage_file
        try:
            leads, info = self.gp.sweep(Http(), "KEY", cap=5, log=lambda m: None)
            self.assertEqual(len(calls), 5)
            self.assertIn("monthly cap", info["stopped"])
            self.assertEqual({l.name for l in leads}, {"Asare Animal Trust", "Ahimsa Animal Care Trust",
                              "CUPA - Compassion Unlimited Plus Action", "Sri Krishna Gaushala"})
            # A second run in the same month gets no more calls.
            calls.clear()
            self.gp.sweep(Http(), "KEY", cap=5, log=lambda m: None)
            self.assertEqual(calls, [])
        finally:
            self.gp.USAGE_FILE = orig

    def test_confirm_only_publishes_web_confirmed_leads(self):
        from scraper import dataset
        known = dataset.base()
        leads = [self.gp.to_lead(p, "q", "Udupi") for p in self.page["places"] if self.gp.keep(p)]
        http = FakeHttp({"lite.duckduckgo": FakeResponse(fixture("ddg_lite_results.html")),
                         "mojeek": FakeResponse(fixture("mojeek_results.html"))})
        s = websearch.Searcher(http, log=lambda m: None, engines=["ddg_lite", "mojeek"])
        confirmed, pending = self.gp.confirm(leads, known, s, log=lambda m: None)
        cupa = [o for o in known if o.awbi_code == "KA004/1993"][0]
        self.assertEqual(cupa.place_id, "ChIJcupa")          # known group: only the place ID added
        asare = [o for o in known if o.name == "Asare Animal Trust"][0]
        self.assertEqual(asare.place_id, "ChIJasare")        # matched your addition
        self.assertEqual([o.name for o in confirmed], [])   # nothing new got confirmed here
        self.assertEqual([l.name for l in pending], ["Sri Krishna Gaushala"])
        for o in known:                                     # no Google address stored anywhere
            self.assertNotIn("Kulai Hosabettu Post, Surathkal, Karnataka 575019, India", o.address)


# ----------------------------------------------------------------- ola maps

class TestOla(unittest.TestCase):
    def setUp(self):
        from scraper.sources import ola_places as ola
        self.ola = ola
        self.orig_usage = ola.USAGE_FILE
        ola.USAGE_FILE = os.path.join(tempfile.mkdtemp(), "ola-usage.json")
        from scraper.net import Cache
        ola._CACHE["c"] = Cache(enabled=False)
        ola.STRICT["on"] = True

    def tearDown(self):
        self.ola.USAGE_FILE = self.orig_usage

    def http(self, text_fixture="ola_textsearch.json"):
        return FakeHttp({"/autocomplete": FakeResponse(fixture(text_fixture)),
                         "/nearbysearch": FakeResponse('{"predictions":[]}'),
                         "/details": FakeResponse(fixture("ola_details.json"))})

    def test_parses_both_response_shapes(self):
        a = self.ola.parse_items(json.loads(fixture("ola_textsearch.json")))
        self.assertEqual(a[0]["name"], "Asare Animal Trust")
        self.assertEqual((a[0]["lat"], a[0]["lng"]), (12.987, 74.79))
        b = self.ola.parse_items(json.loads(fixture("ola_textsearch_alt.json")))
        self.assertEqual(b[0]["name"], "Sri Krishna Gaushala")
        self.assertIn("Kalaburagi", b[0]["address"])
        self.assertEqual(b[0]["id"], "ola-platform:gau9")

    def test_details_normalise_phone_and_social_website(self):
        d = self.ola.parse_details(json.loads(fixture("ola_details.json")))
        self.assertEqual(d["phones"], ["+91 79751 24932"])
        org = self.ola.to_org({"name": "Asare Animal Trust", "district": "Dakshina Kannada",
                               "address": d["address"], "lat": 1, "lng": 2, "website": d["website"],
                               "phones": d["phones"]})
        self.assertEqual(org.website, "")
        self.assertEqual(org.social, ["https://www.facebook.com/asareanimaltrust"])

    def test_sweep_filters_and_reads_details(self):
        orgs, info = self.ola.sweep(self.http(), "KEY", cap=2000, only=["Hassan"], log=lambda m: None, workers=1)
        names = sorted(o.name for o in orgs)
        self.assertEqual(names, ["Asare Animal Trust", "CUPA - Compassion Unlimited Plus Action",
                                 "Hassan Stray Dog Shelter Trust"])
        asare = [o for o in orgs if o.name == "Asare Animal Trust"][0]
        self.assertEqual(asare.district, "Dakshina Kannada")
        self.assertEqual(asare.phones, ["+91 79751 24932"])
        self.assertEqual(asare.sources, ["Ola Maps"])

    def test_every_search_is_biased_to_its_district(self):
        http = self.http()
        self.ola.sweep(http, "KEY", cap=2000, only=["Hassan"], log=lambda m: None, workers=1, with_details=False)
        searches = [p for u, p in zip(http.calls, http.params) if "/autocomplete" in u]
        self.assertEqual(len(searches), len(self.ola.plan(only=["Hassan"])))
        lat, lon = __import__("scraper.geo", fromlist=["x"]).centre("Karnataka", "Hassan")
        centre = [p for p in searches if p["radius"] == "40000"]
        self.assertEqual(len(centre), len(self.ola.CENTRE_TERMS))
        self.assertTrue(all(p["location"] == "{},{}".format(lat, lon) for p in centre))
        points = {p["location"] for p in searches if p["radius"] == "15000"}
        self.assertGreaterEqual(len(points), 2)
        self.assertTrue(all(p["location"] for p in searches))
        self.assertIn("animal trust", [p["input"] for p in centre])

    def test_sweep_stops_at_monthly_cap(self):
        http = self.http()
        orgs, info = self.ola.sweep(http, "KEY", cap=4, only=["Hassan"], log=lambda m: None, workers=1)
        self.assertEqual(len(http.calls), 4)
        self.assertIn("monthly cap", info["stopped"])

    def test_check_reports_bad_key(self):
        http = FakeHttp({"/autocomplete": FakeResponse('{"message":"invalid api key"}', 401)})
        out = []
        self.assertFalse(self.ola.check(http, "BAD", log=out.append))
        self.assertTrue(any("rejected the key" in line for line in out))

    def test_bad_key_is_reported(self):
        http = FakeHttp({"/autocomplete": FakeResponse('{"message":"invalid api key"}', 401)})
        with self.assertRaises(PermissionError):
            self.ola.sweep(http, "BAD", cap=50, only=["Hassan"], log=lambda m: None, workers=1)

    def test_check_command(self):
        out = []
        http = self.http()
        self.assertTrue(self.ola.check(http, "KEY", log=out.append))
        self.assertEqual(len(http.calls), len(self.ola.CHECKS))
        self.assertTrue(http.params[0]["location"])
        self.assertTrue(any("[keep] Asare Animal Trust" in line for line in out))
        self.assertTrue(any("[skip] Happy Tails Pet Clinic" in line for line in out))
        self.assertTrue(http.params[-1].get("place_id"), "details looks up the first place found")

    def test_names_from_the_first_all_india_run(self):
        good = ["Tatanagar GauShala", "Gauri Gopala GauShala", "Asare Animal Trust", "Mysore Snake Rescue",
                "Calcutta Pinjrapole Society Liluah Gosala", "Gau Seva Sadan", "Hope For Animals Rescue Shelter",
                "Fauna On Rescue", "Bombay Panjrapole Cow Shelter Home", "Yahshua Animal Trust"]
        bad = ["India Shelter Home Loans", "Public Night Shelter", "Nagaram Bus Shelter", "Animal Husbandry",
               "Gaurav Krishi Seva Kendra", "Dog Squad Office", "Fire And Rescue Services", "Go Gas Bunk Honakeri",
               "Anamalai Tiger Reserve", "20-41, Near GauShala", "The Force Dog Training Centre",
               "OYO 9316 Green Shelter Guest House", "Shelter Home For Urban Homeless", "Alivira Animal Health"]
        for n in good:
            self.assertTrue(self.ola.name_ok(n), n)
        for n in bad:
            self.assertFalse(self.ola.name_ok(n), n)

    def test_only_animal_groups_are_kept(self):
        item = lambda n: {"name": n, "address": "Karnataka", "types": [], "closed": False}
        for n in ("Asare Animal Trust", "Sri Krishna Gaushala", "Pinjrapole Deesa", "Visakha SPCA"):
            self.assertTrue(self.ola.ola_candidate(item(n)), n)
        for n in ("Asare Charitable Trust", "Mangalore Cafe", "Happy Tails Pet Clinic", "Glamour Burkha & Shawls"):
            self.assertFalse(self.ola.ola_candidate(item(n)), n)

    def test_full_run_publishes_ola_listings_with_credit(self):
        import scraper.__main__ as cli
        fake = FakeHttp({"/autocomplete": FakeResponse(fixture("ola_textsearch.json")),
                         "/details": FakeResponse(fixture("ola_details.json")),
                         "overpass": FakeResponse(fixture("osm_sample.json"))})
        out = tempfile.mkdtemp()
        orig = cli.Http
        cli.Http = lambda **kw: fake
        try:
            code = cli.main(["--out", out, "--site-dir", os.path.join(out, "site"), "--no-cache",
                             "--ola-key", "TEST", "--no-places", "--districts", "Hassan",
                             "--no-web", "--no-enrich", "--no-people", "--no-find-websites",
                             "--awbi-pdf", os.path.join(FIX, "awbi_recognised_ruled.pdf"),
                             "--abc-pdf", os.path.join(FIX, "awbi_abc_ruled.pdf")])
        finally:
            cli.Http = orig
        self.assertEqual(code, 0)
        with open(os.path.join(out, "site", "data", "ngos.json"), encoding="utf-8") as fh:
            site = json.load(fh)
        self.assertTrue(site["meta"]["ola"])
        hassan = [o for o in site["orgs"] if o["n"] == "Hassan Stray Dog Shelter Trust"][0]
        self.assertEqual((hassan["d"], hassan["src"]), ("Hassan", ["Ola Maps"]))
        self.assertTrue(hassan["m"].startswith("https://www.google.com/maps?q="))
        # Ola's listing of a known group merges into it instead of duplicating it.
        self.assertEqual(sum(1 for o in site["orgs"] if "Compassion Unlimited" in o["n"]), 1)
        asare = [o for o in site["orgs"] if o["n"] == "Asare Animal Trust"]
        self.assertEqual(len(asare), 1)
        self.assertIn("Ola Maps", asare[0]["src"])
        import subprocess
        r = subprocess.run([sys.executable, os.path.join(ROOT, "check_site.py"), os.path.join(out, "site")],
                           capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)


# ----------------------------------------------------------- lookup page

class TestLookupPage(unittest.TestCase):
    def test_data_is_embedded_safely(self):
        evil = Org(name='Bad </script><script>alert(1)</script> Trust', source="Web search",
                   district="Udupi", website="javascript:alert(1)", source_urls=["https://example.org/bad"],
                   social=["javascript:alert(2)", "https://instagram.com/ok"],
                   people=[{"key": "x", "name": "Asha Rao", "role": "Founder", "source": "javascript:x"}])
        html = lookup_page.build([evil])
        self.assertEqual(html.count("</script>"), 2)  # only the page's own two script tags
        blob = html.split('<script id="data" type="application/json">')[1].split("</script>")[0]
        data = json.loads(blob.replace("<\\/", "</"))
        rec = data["orgs"][0]
        self.assertEqual(rec["w"], "")
        self.assertEqual(rec["s"], ["https://instagram.com/ok"])
        self.assertEqual(rec["pp"][0]["u"], "")
        self.assertIn("alert(1)", rec["n"])  # kept as text, shown with textContent

    def test_scraped_data_is_tidied_for_the_site(self):
        dash = chr(0x2013)
        a = Org(name="Jeev Daya " + dash + " Gaushala", source="Ola Maps", sources=["Ola Maps"], state="Rajasthan",
                address="Sanganer, Jaipur, Rajasthan 302029", lat="26.8", lon="75.8")
        b = Org(name="Jeev Daya " + dash + " Gaushala", source="Ola Maps", sources=["Ola Maps"], state="Rajasthan",
                address="Sanganer, Jaipur, Rajasthan 302029", lat="26.81", lon="75.8", phones=["+91 99999 00000"])
        nothing = Org(name="Mystery Animal Trust", source="Web search")
        recs = lookup_page.payload([a, b, nothing])["orgs"]
        self.assertEqual([r["n"] for r in recs], ["Jeev Daya - Gaushala"])
        self.assertEqual(recs[0]["p"], ["+91 99999 00000"], "the fuller duplicate is kept")
        self.assertEqual((recs[0]["st"], recs[0]["d"]), ("Rajasthan", "Jaipur"))

    def test_record_fields(self):
        o = Org(name="Captain's Animal Care Trust", source="AWBI recognised list", confidence="high",
                city="Udupi", awbi_code="KA073/2019", awbi_valid_upto="08.01.2030", lat="13.3", lon="74.7",
                people=[{"key": "k", "name": "Asha Rao", "role": "Founder", "source": "https://x.org/team"}])
        r = lookup_page.record(o)
        self.assertEqual((r["d"], r["c"], r["u"]), ("Udupi", ["animal"], False))
        self.assertEqual(r["pp"][0]["h"], "x.org")
        self.assertTrue(r["m"].startswith("https://www.google.com/maps?q="))


# ------------------------------------------------------------- end to end

class TestEndToEnd(unittest.TestCase):
    def test_full_run_writes_clean_spreadsheet(self):
        from openpyxl import load_workbook

        import scraper.__main__ as cli

        routes = {
            "overpass": FakeResponse(fixture("osm_sample.json")),
            "duckduckgo": FakeResponse(fixture("ddg_results.html")),
            "/contact-us": FakeResponse(fixture("site_contact.html"),
                                        url="https://udupistraycare.example.org/contact-us"),
            "/about-us": FakeResponse(fixture("site_about.html"),
                                      url="https://udupistraycare.example.org/about-us"),
            "/our-team": FakeResponse(fixture("site_team.html"),
                                      url="https://udupistraycare.example.org/our-team"),
            "udupistraycare.example.org": FakeResponse(fixture("site_home.html"),
                                                       url="https://udupistraycare.example.org/"),
        }
        fake = FakeHttp(routes)
        original = cli.Http
        cli.Http = lambda **kw: fake
        out = tempfile.mkdtemp()
        try:
            code = cli.main(["--out", out, "--site-dir", os.path.join(out, "site"), "--no-cache",
                             "--no-places", "--districts", "Udupi",
                             "--awbi-pdf", os.path.join(FIX, "awbi_recognised_ruled.pdf"),
                             "--abc-pdf", os.path.join(FIX, "awbi_abc_ruled.pdf")])
        finally:
            cli.Http = original
        self.assertEqual(code, 0)
        wb = load_workbook(os.path.join(out, "animal-ngos.xlsx"))
        ws = wb["NGOs"]
        head = [c.value for c in ws[1]]
        rows = [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
        names = [r["Name"] for r in rows]
        ids = [(r["Name"], r["District"], r["AWBI code"]) for r in rows]
        self.assertEqual(len(ids), len(set(ids)), "duplicates in output")
        # Fixture AWBI rows merge into the built-in snapshot instead of doubling it.
        self.assertEqual(sum(1 for r in rows if r["AWBI code"] == "KA073/2019"), 1)
        self.assertIn("Captain's Animal Care Trust", names)
        self.assertIn("Udupi Stray Care Shelter", names)
        self.assertIn("Mysuru Wildlife Trust", names)
        charlies = [r for r in rows if r["Name"] == "Charlies Animal Rescue Centre (CARE)"][0]
        self.assertIn("OpenStreetMap", charlies["Sources"])
        self.assertEqual(charlies["AWBI code"], "KA067/2018")
        captains = [r for r in rows if r["Name"] == "Captain's Animal Care Trust"][0]
        self.assertIn("AWBI ABC list", captains["Sources"])
        taarini = [r for r in rows if r["Name"] == "Srinivasa Taarini Charitable Trust"][0]
        self.assertEqual((taarini["Category"], taarini["Needs review"]), ("Animal welfare", None))
        shelter = [r for r in rows if r["Name"] == "Udupi Stray Care Shelter"][0]
        self.assertIn("hello@udupistraycare.example.org", shelter["Email"])
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str):
                        self.assertIsNone(LONG_DASH.search(v), v)
        self.assertTrue(os.path.exists(os.path.join(out, "animal-ngos.csv")))
        self.assertIn("Rekha Shenoy (Founder", shelter["Managed by"])
        with open(os.path.join(out, "animal-ngos.html"), encoding="utf-8") as fh:
            page = fh.read()
        self.assertIn("Rekha Shenoy", page)
        self.assertIn("Captain's Animal Care Trust", page)
        self.assertIsNone(LONG_DASH.search(page))
        # The real website folder must never be touched by tests.
        self.assertTrue(os.path.exists(os.path.join(out, "site", "data", "ngos.json")))


class TestEndToEndWithGoogle(unittest.TestCase):
    def test_google_leads_go_to_sheet_not_site(self):
        from openpyxl import load_workbook

        import scraper.__main__ as cli
        from scraper.sources import google_places as gp

        fake = FakeHttp({
            "places.googleapis.com": FakeResponse(fixture("places_page.json")),
            "overpass": FakeResponse(fixture("osm_sample.json")),
        })
        out = tempfile.mkdtemp()
        orig_http, orig_usage = cli.Http, gp.USAGE_FILE
        cli.Http = lambda **kw: fake
        gp.USAGE_FILE = os.path.join(out, "usage.json")
        try:
            code = cli.main(["--out", out, "--site-dir", os.path.join(out, "site"), "--no-cache",
                             "--google-key", "TEST", "--places-max", "3", "--districts", "Udupi",
                             "--no-web", "--no-enrich", "--no-people", "--no-find-websites",
                             "--awbi-pdf", os.path.join(FIX, "awbi_recognised_ruled.pdf"),
                             "--abc-pdf", os.path.join(FIX, "awbi_abc_ruled.pdf")])
        finally:
            cli.Http, gp.USAGE_FILE = orig_http, orig_usage
        self.assertEqual(code, 0)
        self.assertEqual(sum(1 for u in fake.calls if "places.googleapis.com" in u), 3)
        wb = load_workbook(os.path.join(out, "animal-ngos.xlsx"))
        leads = [r[0] for r in wb["Leads to confirm"].iter_rows(min_row=2, values_only=True)]
        self.assertIn("Sri Krishna Gaushala", leads)
        with open(os.path.join(out, "site", "data", "ngos.json"), encoding="utf-8") as fh:
            site = json.load(fh)
        names = [o["n"] for o in site["orgs"]]
        self.assertNotIn("Sri Krishna Gaushala", names)
        asare = [o for o in site["orgs"] if o["n"] == "Asare Animal Trust"][0]
        self.assertEqual(asare["m"], "https://www.google.com/maps/place/?q=place_id:ChIJasare")


class TestDataset(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from scraper import dataset
        cls.orgs = dataset.base()

    def find(self, name):
        return [o for o in self.orgs if o.name == name or name in o.aka]

    def test_size_and_snapshot(self):
        ka = [o for o in self.orgs if o.state == "Karnataka"]
        self.assertGreaterEqual(len(ka), 160)
        self.assertEqual(sum(1 for o in ka if o.awbi_code), 84)
        from scraper import dataset
        if os.path.exists(dataset.NATIONAL_SNAPSHOT):  # saved by the all-India refresh
            self.assertGreater(sum(1 for o in self.orgs if o.awbi_code), 3000)
            self.assertGreaterEqual(len({o.state for o in self.orgs}), 30)

    def test_pfa_bangalore_and_wildlife_hospital_are_one_group(self):
        hits = self.find("PFA Wildlife Hospital")
        self.assertEqual(len(hits), 1)
        self.assertEqual(hits[0].awbi_code, "KA025/2001")
        self.assertIn("wildlife", hits[0].categories)
        self.assertTrue(hits[0].people)

    def test_cupa_lists_suparna_ganguly_with_source(self):
        cupa = self.find("CUPA")[0]
        sg = [p for p in cupa.people if p["name"] == "Suparna Ganguly"]
        self.assertTrue(sg and sg[0]["source"].startswith("https://"))

    def test_no_false_merge_between_similar_names(self):
        wrrc = [o for o in self.orgs if o.awbi_code == "KA017/1999"][0]
        arrt = [o for o in self.orgs if o.awbi_code == "KA006/1987"][0]
        self.assertIn("WRRC", wrrc.aka)
        self.assertNotIn("Wildlife Rescue and Rehabilitation Centre", arrt.aka)

    def test_every_group_is_valid(self):
        from scraper import geo
        for o in self.orgs:
            self.assertIn(o.state, geo.states(), o.name)
            if o.state == "Karnataka":
                self.assertIn(o.district, districts.DISTRICTS, o.name)
            else:
                self.assertIn(o.district, geo.districts_of(o.state) + [""], o.name)
            self.assertTrue(o.source_urls, o.name)
            for u in [o.website] + o.social + o.source_urls:
                self.assertTrue(not u or u.startswith(("http://", "https://")), (o.name, u))

    def test_lookup_payload_search_fields(self):
        data = lookup_page.payload(self.orgs)
        pfa = [r for r in data["orgs"] if r["ac"] == "KA025/2001"][0]
        self.assertIn("PFA Wildlife Hospital", pfa["k"])
        self.assertEqual(data["meta"]["count"], len(self.orgs))


class TestGeo(unittest.TestCase):
    def test_states_and_districts(self):
        from scraper import geo
        self.assertEqual(len(geo.states()), 36)
        self.assertEqual(len(geo.districts_of("Karnataka")), 31)
        self.assertIn("Vijayanagara", geo.districts_of("Karnataka"))
        self.assertGreater(sum(len(geo.districts_of(s)) for s in geo.states()), 700)

    def test_nearest_place_from_coordinates(self):
        from scraper import geo
        self.assertEqual(geo.nearest(12.9651, 74.8139), ("Karnataka", "Dakshina Kannada"))
        self.assertEqual(geo.nearest(26.29, 73.02), ("Rajasthan", "Jodhpur"))
        self.assertEqual(geo.nearest(27.70, 85.32), ("", ""), "Kathmandu is not in India")
        self.assertEqual(geo.nearest("", ""), ("", ""))

    def test_every_district_alias_points_at_a_real_district(self):
        from scraper import geo
        for (st, alias), d in geo.DISTRICT_ALIASES.items():
            self.assertIn(d, geo.districts_of(st), alias)

    def test_normalise_names(self):
        from scraper import geo
        self.assertEqual(geo.normalize_state("Orissa"), "Odisha")
        self.assertEqual(geo.normalize_state("new delhi"), "Delhi")
        self.assertEqual(geo.normalize_state("Atlantis"), "")
        self.assertEqual(geo.normalize_district("Haryana", "Gurgaon"), "Gurugram")
        self.assertEqual(geo.normalize_district("Kerala", "Trivandrum Dist."), "Thiruvananthapuram")

    def test_locate_by_pin_and_name(self):
        from scraper import geo
        self.assertEqual(geo.locate("Kadambody Rd, Surathkal, Mangaluru 575019"), ("Karnataka", "Dakshina Kannada"))
        self.assertEqual(geo.locate("Sector 5, Chandigarh 160005")[0], "Chandigarh")
        st, d = geo.locate("Near bus stand, Kozhikode, Kerala")
        self.assertEqual((st, d), ("Kerala", "Kozhikode"))
        self.assertEqual(geo.locate("Hosapete, Vijayanagara district"), ("Karnataka", "Vijayanagara"))

    def test_org_outside_karnataka_keeps_its_state(self):
        o = Org(name="Animal Rescue Kerala", source="test", state="Kerala", city="Calicut")
        self.assertEqual((o.state, o.district), ("Kerala", "Kozhikode"))
        o = Org(name="Some Trust", source="test", address="12 MG Road, Pune 411001")
        self.assertEqual((o.state, o.district), ("Maharashtra", "Pune"))

    def test_same_name_in_two_states_is_not_merged(self):
        a = Org(name="People For Animals", source="test", state="Kerala", district="Ernakulam", sources=["x"])
        b = Org(name="People For Animals", source="test", state="Odisha", district="Khordha", sources=["y"])
        self.assertEqual(len(merge([a, b])), 2)


class TestNationwide(unittest.TestCase):
    def test_ola_plan_all_india_and_one_state(self):
        from scraper.sources import ola_places
        allp = ola_places.plan()
        self.assertEqual({st for _, st, _, _, _ in allp}, set(__import__("scraper.geo", fromlist=["x"]).states()))
        self.assertTrue(all(loc for _, _, _, loc, _ in allp), "every search needs a location")
        self.assertLess(len(allp), 40000, "must fit the free monthly tier with room for details")
        ka = ola_places.plan(states=["Karnataka"])
        self.assertTrue(ka and all(st == "Karnataka" for _, st, _, _, _ in ka))
        goa = ola_places.plan(states=["Goa"])
        self.assertEqual({d for _, _, d, _, _ in goa}, {"North Goa", "South Goa"})
        self.assertIn(("animal trust", "Goa", "North Goa"), [x[:3] for x in goa])

    def test_site_payload_has_states_and_map(self):
        orgs = [Org(name="Goa Animal Welfare Trust", source="test", state="Goa", district="North Goa", sources=["AWBI"],
                    source_urls=["https://example.org"]),
                Org(name="Udupi SPCA", source="test", district="Udupi", sources=["AWBI"], source_urls=["https://example.org"])]
        data = lookup_page.payload(orgs)
        self.assertEqual({r["st"] for r in data["orgs"]}, {"Goa", "Karnataka"})
        self.assertEqual(len(data["meta"]["states"]), 36)
        goa = dict((d[0], d[1:]) for d in data["meta"]["states"]["Goa"])
        self.assertIsNotNone(goa["North Goa"][0])
        self.assertNotIn("shapes", data["meta"], "outlines ship as their own file")
        data["meta"].update(lookup_page.state_shapes())
        shapes = data["meta"]["shapes"]
        self.assertEqual(set(shapes), set(__import__("scraper.geo", fromlist=["x"]).states()) - {"Ladakh"})
        self.assertIn("MIT", data["meta"]["map_credit"])
        for rings in shapes.values():
            for r in rings:
                self.assertEqual(len(r) % 2, 0)
                self.assertTrue(all(6 < r[i + 1] < 38 and 67 < r[i] < 98 for i in range(0, len(r), 2)))



class TestRapidResponseEmbed(unittest.TestCase):
    def test_every_rule_is_scoped_to_the_container(self):
        from scraper import rrcss
        css = rrcss.build()
        body = re.sub(r"/\*.*?\*/", "", css, flags=re.S)
        body = re.sub(r"@keyframes[^{]*\{([^{}]*\{[^{}]*\})*[^{}]*\}", "", body)
        for sel in re.findall(r"([^{}]+)\{[^{}]*\}", body):
            for part in sel.replace("@media", "").split(","):
                part = part.strip()
                if part.startswith("(") or not part:
                    continue
                self.assertTrue(part.startswith(".ngo-rr"), part)

    def test_scoping(self):
        from scraper import rrcss
        out = rrcss.scope(":root{--a:1}html,body{margin:0}.a,.b:not(.c,.d){x:y}@media (max-width:9px){h1{z:1}}@keyframes k{to{o:1}}")
        self.assertIn(".ngo-rr{--a:1}", out)
        self.assertIn(".ngo-rr{margin:0}", out)
        self.assertIn(".ngo-rr .a,.ngo-rr .b:not(.c,.d){x:y}", out)
        self.assertIn("@media (max-width:9px){.ngo-rr h1{z:1}}", out)
        self.assertIn("@keyframes k{to{o:1}}", out)

    def test_site_ships_the_themed_sheet_and_loads_from_beside_the_script(self):
        out = tempfile.mkdtemp()
        lookup_page.copy_static(out)
        self.assertTrue(os.path.exists(os.path.join(out, "rr.css")))
        with open(os.path.join(ROOT, "scraper", "web", "app.js"), encoding="utf-8") as fh:
            js = fh.read()
        self.assertNotRegex(js, r'fetch\("data/')
        self.assertNotIn('import("./', js)
        for f in ("common.js", "public.js"):
            with open(os.path.join(ROOT, "scraper", "web", "fed", f), encoding="utf-8") as fh:
                self.assertNotRegex(fh.read(), r'(fetch\("data/|href: "join\.html)')

    def test_embedded_markup_has_every_icon_and_control(self):
        with open(os.path.join(ROOT, "scraper", "web", "index.html"), encoding="utf-8") as fh:
            page = fh.read()
        with open(os.path.join(ROOT, "scraper", "web", "app.js"), encoding="utf-8") as fh:
            js = fh.read()
        for sym in re.findall(r'<symbol id="([^"]+)"', page):
            self.assertIn('id="%s"' % sym, js, sym)
        for i in ("types", "searchbox", "q", "clear", "crumbs", "title", "sub", "filterBtn", "filterDot", "share",
                  "filters", "fContact", "fPeople", "fAwbi", "fVerified", "fFed", "view", "foot", "toast", "needsBtn"):
            self.assertIn('id="%s"' % i, js, i)


class TestRrtRules(unittest.TestCase):
    def test_block_goes_into_rrt_rules_once(self):
        sys.path.insert(0, os.path.join(ROOT, "tools"))
        import rrt_rules
        with open(os.path.join(ROOT, "tests", "rules", "rrt-base.rules"), encoding="utf-8") as fh:
            base = fh.read()
        once = rrt_rules.merge(base)
        self.assertEqual(once, rrt_rules.merge(once))
        self.assertEqual(once.count("NGO-DIRECTORY START"), 1)
        self.assertEqual(once.count("{"), once.count("}"))
        self.assertTrue(once.startswith(base[:base.rindex("}", 0, base.rindex("}"))].rstrip()))
        block = rrt_rules.block()
        self.assertIn("match /ngo/dir {", block)
        # every helper in the block is the directory's own, so nothing clashes with RRT's
        for name in re.findall(r"function (\w+)\(", block):
            self.assertTrue(name.startswith("ngo"), name)
        for path in re.findall(r"/databases/\$\(database\)/documents/([a-z_]+)/", block):
            self.assertIn(path, ("ngo", "admins"), path)


class TestRemoved(unittest.TestCase):
    RULES = [{"name_contains": "sudheendra aithal", "state": "Karnataka", "note": "x"}]

    def test_taken_off_by_name_other_name_or_person_in_that_state(self):
        from scraper import removed
        base = {"n": "Animal Rescue Center By Sudheendra Aithal", "k": [], "pp": [], "st": "Karnataka"}
        self.assertTrue(removed.matches(base, self.RULES))
        self.assertTrue(removed.matches(dict(base, n="Saligrama Rescue", k=["Sudheendra Aithal's shelter"]), self.RULES))
        self.assertTrue(removed.matches(dict(base, n="Saligrama Rescue", pp=[{"n": "Sudheendra Aithal"}]), self.RULES))
        self.assertFalse(removed.matches(dict(base, st="Kerala"), self.RULES))
        self.assertFalse(removed.matches(dict(base, n="Sudheendra Aithalabad Trust"), self.RULES))

    def test_builds_and_the_live_data_leave_them_out(self):
        from scraper import removed
        self.assertTrue(removed.load())
        with open(os.path.join(ROOT, "site", "data", "ngos.json"), encoding="utf-8") as fh:
            orgs = json.load(fh)["orgs"]
        self.assertFalse([o["n"] for o in orgs if removed.matches(o)])
        kept = removed.keep([{"n": "Animal Rescue Center By Sudheendra Aithal", "k": [], "pp": [], "st": "Karnataka"},
                             {"n": "CUPA", "k": [], "pp": [], "st": "Karnataka"}])
        self.assertEqual([o["n"] for o in kept], ["CUPA"])


class TestFederation(unittest.TestCase):
    def test_site_build_ships_federation_files_but_not_the_mock(self):
        out = tempfile.mkdtemp()
        lookup_page.copy_static(out)
        for f in ("join.html", "admin.html", "fed/api.js", "fed/join.js", "fed/admin.js", "fed/public.js", "vendor/firebase.js"):
            self.assertTrue(os.path.exists(os.path.join(out, f)), f)
        self.assertFalse(os.path.exists(os.path.join(out, "fed", "mock.js")))

    def test_money_check_is_the_same_in_the_browser_and_the_rules(self):
        with open(os.path.join(ROOT, "firestore.rules"), encoding="utf-8") as fh:
            rules = fh.read()
        with open(os.path.join(ROOT, "scraper", "web", "fed", "common.js"), encoding="utf-8") as fh:
            js = fh.read()
        rule_re = re.search(r"s\.matches\('\(\?is\)\.\*\((.*)\)\.\*'\)", rules).group(1).replace("\\\\", "\\")
        js_re = re.search(r'new RegExp\("\(" \+ String\.fromCharCode\(0x20b9\) \+ "\|(.*)\)", "is"\)', js).group(1).replace("\\\\", "\\")
        self.assertEqual(rule_re.split("|", 1)[1], js_re)

    def test_money_patterns(self):
        with open(os.path.join(ROOT, "scraper", "web", "fed", "common.js"), encoding="utf-8") as fh:
            js = fh.read()
        body = re.search(r'new RegExp\("\(" \+ String\.fromCharCode\(0x20b9\) \+ "\|(.*)\)", "is"\)', js).group(1).replace("\\\\", "\\")
        money = re.compile("(" + chr(0x20b9) + "|" + body + ")", re.I | re.S)
        for bad in ("Send to UPI gau@okaxis", "Bank account no 1234", "Need Rs 5000", "Need " + chr(0x20b9) + "500",
                    "Scan the QR code", "Fundraising drive", "GPay 98765", "Donations in cash", "G-Pay 98765",
                    "Phone-Pe", "pay to cows@ibl", "Sponsor a cow", "1500/- per month", "NEFT details", "U P I",
                    "PayPal me", "Milaap campaign", "acct no 1234"):
            self.assertTrue(money.search(bad), bad)
        for ok in ("We need blankets and volunteers, please donate old bedsheets", "Dry fodder, 2 tonnes",
                   "Foster homes for 3 puppies", "Vet help for a cow with a fracture", "Near SBI Bank, Bank Colony",
                   "Cow shelter behind the bus stand"):
            self.assertFalse(money.search(ok), ok)

    def test_rules_have_the_verification_gates(self):
        with open(os.path.join(ROOT, "firestore.rules"), encoding="utf-8") as fh:
            rules = fh.read()
        for must in ('ngoApp().get("docsOk", false) == true', "ngoApp().vouch.ok == true", "ngoApp().vouch.by != ngoApp().uid",
                     "ngoApp().vouch.by != request.auth.uid", 'ngoApp().visit.ok == true', 'ngoAudited("vouch", id)',
                     'ngoApp().vouch.role == "coordinator" || ngoApp().get("needsVisit", false) == true',
                     'ngoNextComment().status == "pending"', 'allow list: if resource.data.status == "approved" || ngoIsAdmin();',
                     'id.matches(".*_" + request.auth.uid)', 'allow update, delete: if false;'):
            self.assertIn(must, rules)


class TestNetworks(unittest.TestCase):
    def test_people_for_animals_units_name_the_national_network(self):
        from scraper import networks
        rec = networks.apply({"n": "PFA Agra", "k": [], "nt": [], "sr": []})
        self.assertTrue(any("Maneka Gandhi" in n for n in rec["nt"]))
        self.assertIn("https://en.wikipedia.org/wiki/People_for_Animals", rec["sr"])
        other = networks.apply({"n": "Gopal Das Gau Sewa Ashram", "k": [], "nt": [], "sr": []})
        self.assertEqual(other["nt"], [])
        again = networks.apply(rec)
        self.assertEqual(sum("Maneka Gandhi" in n for n in again["nt"]), 1)

class TestNoLongDashesInProject(unittest.TestCase):
    def test_source_files(self):
        for base, _, files in os.walk(ROOT):
            if ".venv" in base or "__pycache__" in base:
                continue
            for f in files:
                if f.endswith((".py", ".md", ".sh", ".txt", ".html", ".json", ".js", ".css", ".csv", ".svg")):
                    with open(os.path.join(base, f), encoding="utf-8") as fh:
                        text = fh.read()
                    self.assertIsNone(LONG_DASH.search(text), os.path.join(base, f))


if __name__ == "__main__":
    unittest.main()

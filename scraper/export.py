"""Write the results to Excel and CSV."""

import csv

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .classify import categories, label
from .people import label as people_label

COLUMNS = [
    ("Name", 38), ("Category", 26), ("Managed by", 40), ("District", 18), ("City / area", 16),
    ("Address", 40), ("Phone", 22), ("Email", 30), ("Website", 32),
    ("Social", 32), ("AWBI code", 12), ("AWBI valid upto", 14),
    ("Sources", 26), ("Confidence", 11), ("Needs review", 12),
    ("Map", 20), ("Notes", 50), ("Source links", 40), ("Managed by: where found", 40),
]


def row_for(org):
    cats = categories(org.name, org.snippet)
    if not cats and any(s.startswith("AWBI") for s in org.sources):
        cats = ["animal"]  # AWBI only recognises animal welfare bodies
    needs_review = "Yes" if org.confidence == "low" or not cats else ""
    maps = ""
    if org.lat and org.lon:
        maps = "https://www.google.com/maps?q={},{}".format(org.lat, org.lon)
    return [
        org.name, label(cats), people_label(org), org.district, org.city, org.address,
        "; ".join(org.phones), "; ".join(org.emails), org.website,
        "; ".join(org.social), org.awbi_code, org.awbi_valid_upto,
        ", ".join(org.sources), org.confidence.title(), needs_review, maps,
        "; ".join(org.notes), "; ".join(org.source_urls),
        "; ".join(sorted({p["source"] for p in org.people})),
    ]


def sort_key(org):
    rank = {"high": 0, "medium": 1, "low": 2}[org.confidence]
    return (org.district or "zzz", rank, org.name.lower())


def write(orgs, xlsx_path, csv_path, run_log, leads=()):
    orgs = sorted(orgs, key=sort_key)
    rows = [row_for(o) for o in orgs]

    wb = Workbook()
    ws = wb.active
    ws.title = "NGOs"
    ws.append([c for c, _ in COLUMNS])
    for r in rows:
        ws.append(r)
    head_fill = PatternFill("solid", fgColor="1F4E46")
    for i, (_, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:{}{}".format(get_column_letter(len(COLUMNS)), len(rows) + 1)

    by_district = {}
    for o in orgs:
        by_district[o.district or "Unknown"] = by_district.get(o.district or "Unknown", 0) + 1
    ds = wb.create_sheet("By district")
    ds.append(["District", "Organisations"])
    for d, n in sorted(by_district.items(), key=lambda x: -x[1]):
        ds.append([d, n])
    ds.column_dimensions["A"].width = 30
    ds["A1"].font = ds["B1"].font = Font(bold=True)

    if leads:
        ls = wb.create_sheet("Leads to confirm")
        ls.append(["Name on Google Maps", "District", "Open in Google Maps", "Found by searching",
                   "What to do"])
        for c in ls[1]:
            c.font = Font(bold=True)
        for lead in sorted(leads, key=lambda l: (l.district, l.name.lower())):
            ls.append([lead.name, lead.district, lead.maps_link, lead.queries[0] if lead.queries else "",
                       "Call to confirm it is an active NGO, then add it to scraper/data/my_additions.csv"])
        for col, w in zip("ABCDE", (40, 20, 58, 44, 70)):
            ls.column_dimensions[col].width = w
        ls.freeze_panes = "A2"

    log = wb.create_sheet("Run log")
    log.append(["Item", "Detail"])
    log["A1"].font = log["B1"].font = Font(bold=True)
    for k, v in run_log:
        log.append([k, v])
    log.column_dimensions["A"].width = 34
    log.column_dimensions["B"].width = 120
    wb.save(xlsx_path)

    with open(csv_path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow([c for c, _ in COLUMNS])
        w.writerows(rows)
    return len(rows)
